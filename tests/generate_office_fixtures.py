import pathlib

from docx import Document
from pptx import Presentation

OUTPUT = pathlib.Path(__file__).with_name("fixtures")
OUTPUT.mkdir(exist_ok=True)

document = Document()
document.add_heading("Classroom Downloader", level=1)
document.add_paragraph("Gotenberg DOCX integration fixture")
document.save(OUTPUT / "sample.docx")

presentation = Presentation()
slide = presentation.slides.add_slide(presentation.slide_layouts[1])
slide.shapes.title.text = "Classroom Downloader"
slide.placeholders[1].text = "Gotenberg PPTX integration fixture"
presentation.save(OUTPUT / "sample.pptx")
